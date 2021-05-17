from openpyxl import Workbook
# from openpyxl.compat import range
from openpyxl.utils import get_column_letter
from openpyxl.chart import (
    AreaChart3D,
    AreaChart,
    Reference,
    Series,
)

wb = Workbook()
ws = wb.active

rows = [
    ['Number', 'Batch 1', 'Batch 2'],
    ['A', 40, 60],
    ['B', 40, 75],
    ['C', 50, 60],
    ['D', 30, 50],
    ['E', 25, 65],
    ['F', 50, 70],
]

for row in rows:
    ws.append(row)

chart = AreaChart3D()
chart.title = "Area Chart"
chart.style = 13
chart.x_axis.title = 'Test'
chart.y_axis.title = 'Percentage'

# yData needs to start on row=1 to include titles
xData = Reference(ws, min_col=1, min_row=2, max_row=7)
yData = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=7)
chart.add_data(yData, titles_from_data=True)
chart.set_categories(xData)
ws.add_chart(chart, "B10")

wb.save('data/charts3D.xlsx')