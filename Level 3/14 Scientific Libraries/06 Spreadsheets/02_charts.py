from openpyxl import Workbook
#from openpyxl.compat import range
from openpyxl.utils import get_column_letter
from openpyxl.chart import (
    AreaChart3D,
    AreaChart,
    Reference,
    Series,
)

wb = Workbook()
ws = wb.active

rows = [
    ['Number', 'Batch 1', 'Batch 2'],
    ['A', 40, 30],
    ['B', 40, 625],
    ['C', 50, 30],
    ['D', 30, 100],
    ['E', 25, 555],
    ['F', 50, 10],
]

for row in rows:
    ws.append(row)

chart = AreaChart()
chart.title = "Area Chart"
chart.style = 13
chart.x_axis.title = 'Test'
chart.y_axis.title = 'Percentage'

# yData needs to start on row=1 to include titles
xData = Reference(ws, min_col=1, min_row=2, max_row=7)
yData = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=7)
chart.add_data(yData, titles_from_data=True)
chart.set_categories(xData)
ws.add_chart(chart, "C10")


chart = AreaChart3D()
cats = Reference(ws, min_col=1, min_row=1, max_row=7)
data = Reference(ws, min_col=2, min_row=1, max_col=3, max_row=7)
chart.add_data(data, titles_from_data=True)
# chart.set_categories(cats)
ws.add_chart(chart, "B30")



wb.save('data/charts2.xlsx')